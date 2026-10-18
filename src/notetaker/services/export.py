"""Offline document rendering. Production exports require a content-bound review."""
import os
import tempfile
from importlib.resources import files
from pathlib import Path
from xml.sax.saxutils import escape
from notetaker.domain.minutes import require_review

DISCLAIMER = "BROUILLON NON OFFICIEL - contrôle humain requis avant utilisation"
TITLES = {"summary": "Compte rendu synthétique", "transcript": "Transcription intégrale", "municipal": "Projet de procès-verbal municipal"}


def export_draft(content, review, destination: Path, *, session_id: str, municipality: str, session_date: str, votes=None, decisions=None, details=None, reviewer_role=None):
    require_review(content, review)
    if content.get("document_mode") == "municipal" and reviewer_role != "secretary":
        raise ValueError("MUNICIPAL_SECRETARY_REVIEW_REQUIRED")
    import json
    if "\ufffd" in json.dumps([content, municipality, session_date, votes, decisions, details], ensure_ascii=False):
        raise ValueError("DAMAGED_TRANSCRIPT_RETRANSCRIBE_ORIGINAL")
    from .storage import validate_votes
    votes, decisions = votes or [], decisions or []
    validate_votes(votes)
    suffix = destination.suffix.lower()
    if suffix not in {".docx", ".pdf"}:
        raise ValueError("UNSUPPORTED_EXPORT_FORMAT")
    if any(p["vote_id"] is not None and p["vote_id"] not in {v["id"] for v in votes} for p in content["points"]):
        raise ValueError("UNKNOWN_STRUCTURED_VOTE")
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(dir=destination.parent, suffix=suffix)
    os.close(fd)
    temporary = Path(name)
    metadata = " | ".join(x for x in (municipality, session_date) if x)
    try:
        (_docx if suffix == ".docx" else _pdf)(content, metadata, temporary, votes, decisions, details or {})
        with temporary.open("r+b") as stream:
            os.fsync(stream.fileno())
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)
    return destination


def vote_lines(votes):
    for vote in votes:
        if vote["status"] == "unknown":
            yield f"{vote['title']} : résultat non renseigné"
        else:
            yield (f"{vote['title']} : pour {vote['for_count']}, contre {vote['against_count']}, "
                   f"abstentions {vote['abstain_count']}, non-participations {vote['nonparticipating_count']}. Source saisie : {vote['source']}")


def document_sections(content, votes, decisions, details):
    municipal = content.get("document_mode") == "municipal"
    if municipal:
        yield "Cadre de la séance", "\n".join(f"{label} : {details.get(key) or 'À compléter'}" for key, label in
            [("date", "Date"), ("location", "Lieu"), ("participants", "Participants et présences"), ("secretary", "Secrétaire de séance")])
    for point in content["points"]:
        yield point["title"], point["summary"]
    if decisions:
        yield "Décisions confirmées par l’opérateur", "\n\n".join(decisions)
    if votes:
        yield "Scrutins vérifiés par l’opérateur", "\n\n".join(vote_lines(votes))
    if municipal and not decisions:
        yield "Décisions et délibérations", "Aucune décision confirmée saisie. À compléter après vérification ; aucune adoption n’est déduite du débat."
    if municipal and not votes:
        yield "Scrutins", "Aucun résultat vérifié saisi. L’absence de saisie ne prouve pas l’absence de scrutin."
    if content["to_verify"]:
        yield "Points de contrôle avant utilisation", "\n\n".join(content["to_verify"])


def _docx(content, metadata, path, votes, decisions, details=None):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    details = details or {}
    sections = list(document_sections(content, votes, decisions, details))
    validate_glyphs([TITLES[content.get("document_mode", "summary")], metadata,
                     details.get("title", "Séance"), DISCLAIMER, *(h + t for h, t in sections)])
    doc = Document()
    from .font_embedding import embed_docx_fonts
    embed_docx_fonts(doc)
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.3)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Noto Sans"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    section.header.paragraphs[0].text = metadata[:100]
    section.different_first_page_header_footer = True
    title = TITLES[content.get("document_mode", "summary")]
    doc.add_paragraph("NOTETAKER")
    doc.add_paragraph(title, "Title")
    doc.add_paragraph(details.get("title", "Séance"))
    doc.add_paragraph(metadata)
    doc.add_paragraph(DISCLAIMER)
    doc.add_paragraph("Document de travail. La relecture ne vaut pas adoption officielle.")
    if details.get("preview"):
        doc.add_paragraph("APERÇU TECHNIQUE NON VALIDÉ — aucune relecture humaine enregistrée")
    doc.add_page_break()
    for heading, text in sections:
        if content.get("document_mode") == "transcript" and heading != "Points de contrôle avant utilisation":
            paragraph = doc.add_paragraph()
            paragraph.add_run(heading + "\n").bold = True
            paragraph.add_run(text)
        else:
            doc.add_heading(heading, 1)
            for paragraph in text.split("\n\n"):
                doc.add_paragraph(paragraph)
    footer = section.footer.paragraphs[0]
    footer.add_run("NoteTaker | Projet | Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.save(path)


def pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    folder = files("notetaker").joinpath("assets/fonts")
    if "NoteTakerNoto" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("NoteTakerNoto", str(folder.joinpath("NotoSans-Regular.ttf"))))
        pdfmetrics.registerFont(TTFont("NoteTakerNotoBold", str(folder.joinpath("NotoSans-Bold.ttf"))))
        pdfmetrics.registerFontFamily("NoteTakerNoto", normal="NoteTakerNoto", bold="NoteTakerNotoBold", italic="NoteTakerNoto", boldItalic="NoteTakerNotoBold")
    return pdfmetrics.getFont("NoteTakerNoto")


def validate_glyphs(texts):
    font = pdf_fonts()
    missing = sorted({ord(c) for text in texts for c in text if not c.isspace() and ord(c) not in font.face.charToGlyph})
    if missing:
        raise ValueError("UNSUPPORTED_EXPORT_GLYPHS " + ", ".join(f"U+{c:04X}" for c in missing))


def _pdf(content, metadata, path, votes, decisions, details=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    details = details or {}
    title = TITLES[content.get("document_mode", "summary")]
    sections = list(document_sections(content, votes, decisions, details))
    texts = [title, metadata, details.get("title", "Séance"), DISCLAIMER, *(h + t for h, t in sections)]
    validate_glyphs(texts)
    body = ParagraphStyle("Body", fontName="NoteTakerNoto", fontSize=10, leading=15, spaceAfter=9, textColor=colors.HexColor("#202b35"))
    heading = ParagraphStyle("Heading", parent=body, fontName="NoteTakerNotoBold", fontSize=13, leading=18, spaceBefore=15, spaceAfter=9, keepWithNext=True)
    cover = ParagraphStyle("Cover", parent=heading, fontSize=27, leading=35, spaceAfter=24)
    small = ParagraphStyle("Small", parent=body, fontSize=9, leading=13)
    def para(text, style=body):
        return Paragraph(escape(text).replace("\n", "<br/>"), style)
    flow = [Spacer(1, 70), para("NOTETAKER", small), Spacer(1, 20), para(title, cover),
            para(details.get("title", "Séance"), heading), para(metadata), Spacer(1, 28), para(DISCLAIMER, small),
            para("Document de travail. La relecture ne vaut pas adoption officielle.", small)]
    if details.get("preview"):
        flow += [Spacer(1, 20), para("APERÇU TECHNIQUE NON VALIDÉ — aucune relecture humaine enregistrée", heading)]
    flow.append(PageBreak())
    for h, text in sections:
        if content.get("document_mode") == "transcript" and h != "Points de contrôle avant utilisation":
            flow.append(Paragraph("<b>" + escape(h) + "</b><br/>" + escape(text).replace("\n", "<br/>"), body))
        else:
            flow.append(para(h, heading))
            flow += [para(p) for p in text.split("\n\n")]
    def page(canvas, doc):
        canvas.saveState()
        canvas.setFont("NoteTakerNoto", 8)
        if doc.page > 1:
            canvas.setStrokeColor(colors.HexColor("#cbd2d8"))
            canvas.line(54, A4[1] - 47, A4[0] - 54, A4[1] - 47)
            canvas.drawString(54, A4[1] - 36, metadata[:90])
        canvas.drawString(54, 30, "NoteTaker | " + ("APERÇU NON VALIDÉ" if details.get("preview") else "Projet non officiel"))
        canvas.drawRightString(A4[0] - 54, 30, f"Page {doc.page}")
        canvas.restoreState()
    SimpleDocTemplate(str(path), pagesize=A4, rightMargin=54, leftMargin=54, topMargin=66, bottomMargin=55,
                      title=title, author="NoteTaker").build(flow, onFirstPage=page, onLaterPages=page)
