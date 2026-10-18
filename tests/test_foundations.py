import hashlib
import json
import os
import sqlite3
import subprocess
import sys
import wave
from pathlib import Path
import pytest
from jsonschema import ValidationError
from notetaker.domain.minutes import validate_minutes, record_human_review
from notetaker.engines.router import run_with_fallback, EngineUnavailable
from notetaker.infrastructure.atomic import atomic_json, sha256_file
from notetaker.infrastructure.database import open_database, migrate, recover_jobs, backup_database
from notetaker.infrastructure.models import verify_assets
from notetaker.infrastructure.offline import enforce_offline, local_path, OFFLINE_ENV
from notetaker.infrastructure.process import run_worker
from notetaker.services.audio_segments import commit_pcm_segment, inspect_recovery
from notetaker.services.export import export_draft

@pytest.fixture
def draft():
    return {"points": [{"title": "Travaux de la bibliothèque", "summary": "Le devis doit être vérifié.", "sources": ["s0"], "vote_id": None}], "to_verify": ["Résultat du vote à compléter"]}

def test_offline_policy_overrides_host_preferences(monkeypatch):
    for key in OFFLINE_ENV:
        monkeypatch.setenv(key, "unsafe-value")
    enforce_offline()
    assert all(os.environ[k] == v for k, v in OFFLINE_ENV.items())

def test_hub_id_does_not_trigger_a_download():
    with pytest.raises(FileNotFoundError):
        local_path("Systran/faster-whisper-small", directory=True)

def test_missing_or_changed_models_fail(tmp_path):
    asset = tmp_path / "model.bin"
    asset.write_bytes(b"unit-test-integrity-payload")
    entry = {"path": asset.name, "size_bytes": asset.stat().st_size, "sha256": sha256_file(asset)}
    verify_assets(tmp_path, [entry])
    with pytest.raises(ValueError, match="REDISTRIBUTION"):
        verify_assets(tmp_path, [entry], release=True)
    asset.write_bytes(b"modified")
    with pytest.raises(ValueError, match="INTEGRITY"):
        verify_assets(tmp_path, [entry])
    asset.unlink()
    with pytest.raises(FileNotFoundError):
        verify_assets(tmp_path, [entry])

def test_manifest_traversal_rejected(tmp_path):
    with pytest.raises(ValueError, match="UNSAFE"):
        verify_assets(tmp_path, [{"path": "../outside.bin"}])

def test_empty_manifest_is_not_a_valid_release(tmp_path):
    with pytest.raises(ValueError, match="EMPTY"):
        verify_assets(tmp_path, [], release=True)

def test_gpu_failure_retries_cpu_in_order():
    # Fault injection for routing only; this is not an inference test.
    devices = []
    def failing_gpu(device):
        devices.append(device)
        if device == "cuda":
            raise RuntimeError("injected GPU failure")
        return {"completed": True}
    result = run_with_fallback(failing_gpu)
    assert devices == ["cuda", "cpu"]
    assert result.device == "cpu" and not result.attempts[0].ok

def test_cpu_mode_never_probes_gpu():
    devices = []
    run_with_fallback(lambda d: devices.append(d) or {}, "cpu")
    assert devices == ["cpu"]

def test_no_working_engine_is_an_error():
    def fail(device):
        raise OSError("injected missing DLL")
    with pytest.raises(EngineUnavailable) as caught:
        run_with_fallback(fail)
    assert len(caught.value.attempts) == 2

def test_cancellation_does_not_start_cpu():
    def cancel(device):
        raise KeyboardInterrupt()
    with pytest.raises(KeyboardInterrupt):
        run_with_fallback(cancel)

def test_sources_and_votes_are_constrained(draft):
    validate_minutes(draft, {"s0"}, {})
    with pytest.raises(ValueError, match="UNKNOWN_SOURCE"):
        validate_minutes(draft, set(), {})
    draft["points"][0]["vote_id"] = "invented-vote"
    with pytest.raises(ValueError, match="UNKNOWN_STRUCTURED_VOTE"):
        validate_minutes(draft, {"s0"}, {})

def test_model_cannot_approve_or_invent_numeric_vote_fields(draft):
    draft["approved"] = True
    with pytest.raises(ValidationError):
        validate_minutes(draft, {"s0"}, {})
    del draft["approved"]
    draft["points"][0]["for_count"] = 21
    with pytest.raises(ValidationError):
        validate_minutes(draft, {"s0"}, {})

def test_exports_require_review_bound_to_content(tmp_path, draft):
    kwargs = {"session_id": "test-0", "municipality": "Commune de test", "session_date": "2026-10-16"}
    dest = tmp_path / "draft.pdf"
    with pytest.raises(ValueError, match="HUMAN_REVIEW"):
        export_draft(draft, None, dest, **kwargs)
    assert not dest.exists()
    review = record_human_review(draft, "Test operator")
    draft["points"][0]["summary"] = "Révision après relecture"
    with pytest.raises(ValueError, match="HUMAN_REVIEW"):
        export_draft(draft, review, dest, **kwargs)

@pytest.mark.parametrize("suffix", ["docx", "pdf"])
def test_real_local_exports(tmp_path, draft, suffix):
    review = record_human_review(draft, "Automated test fixture, not product approval")
    path = export_draft(draft, review, tmp_path / f"draft.{suffix}", session_id="test-0", municipality="Commune de test", session_date="2026-10-16")
    if suffix == "docx":
        from docx import Document
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Commune de test" in doc.sections[0].header.paragraphs[0].text
        assert "test-0" not in text  # Technical identifiers stay in the internal snapshot.
    else:
        from pypdf import PdfReader
        reader = PdfReader(path)
        assert len(reader.pages) == 2  # Cover plus body.
        text = "\n".join(p.extract_text() for p in reader.pages)
        assert "Page 1" in text
    assert "BROUILLON NON OFFICIEL" in text and "à compléter" in text

def test_database_restart_and_recovery(tmp_path):
    path = tmp_path / "sessions.sqlite3"
    connection = open_database(path)
    with connection:
        connection.execute("INSERT INTO session(id,title) VALUES('s0','Test')")
        connection.execute("INSERT INTO job VALUES('j0','s0','asr','running',NULL)")
        connection.execute("INSERT INTO job VALUES('j1','s0','asr','completed',NULL)")
    connection.close()
    connection = open_database(path)
    assert recover_jobs(connection) == 1
    assert recover_jobs(connection) == 0
    assert connection.execute("SELECT status FROM job WHERE id='j1'").fetchone()[0] == "completed"
    assert connection.execute("PRAGMA foreign_keys").fetchone()[0] == 1
    connection.close()

def test_failed_migration_rolls_back_every_statement(tmp_path):
    connection = open_database(tmp_path / "test.sqlite3")
    with pytest.raises(sqlite3.OperationalError):
        migrate(connection, {1: [], 2: [], 3: [], 4: [], 5: ["CREATE TABLE should_rollback(id TEXT)", "INVALID SQL"]})
    assert connection.execute("PRAGMA user_version").fetchone()[0] == 4
    assert connection.execute("SELECT name FROM sqlite_master WHERE name='should_rollback'").fetchall() == []
    connection.close()

def test_sqlite_backup_includes_committed_wal(tmp_path):
    connection = open_database(tmp_path / "source.sqlite3")
    with connection:
        connection.execute("INSERT INTO session(id,title) VALUES('s0','Test')")
    backup = tmp_path / "backup.sqlite3"
    backup_database(connection, backup)
    with sqlite3.connect(backup) as restored:
        assert restored.execute("SELECT title FROM session").fetchone()[0] == "Test"
        assert restored.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    connection.close()

def test_segment_survives_a_real_child_process_crash(tmp_path):
    code = """
import os,sys
from pathlib import Path
from notetaker.services.audio_segments import commit_pcm_segment
directory=Path(sys.argv[1])
commit_pcm_segment(directory, b'\\0\\0' * 16000)
with (directory/'interrupted.partial').open('wb') as f:
    f.write(b'incomplete PCM fixture'); f.flush(); os.fsync(f.fileno())
os._exit(23)
"""
    env = dict(os.environ, PYTHONPATH=str(Path(__file__).resolve().parents[1] / "src"))
    result = subprocess.run([sys.executable, "-c", code, str(tmp_path)], env=env)
    assert result.returncode == 23
    recovery = inspect_recovery(tmp_path)
    assert len(recovery["valid"]) == 1 and not recovery["damaged"]
    assert recovery["partial"] == ["interrupted.partial"]
    with wave.open(str(tmp_path / recovery["valid"][0])) as audio:
        assert audio.getnframes() == 16000

def test_orphan_and_corrupt_segments_are_reported_without_deletion(tmp_path):
    path = commit_pcm_segment(tmp_path, b"\0\0" * 100)
    path.write_bytes(b"corrupt")
    orphan = tmp_path / "orphan.wav"
    orphan.write_bytes(b"orphan fixture")
    recovery = inspect_recovery(tmp_path)
    assert recovery["damaged"] == [path.name]
    assert recovery["orphans"] == [orphan.name]
    assert path.exists() and orphan.exists()

def test_worker_missing_model_fails_without_download(tmp_path):
    request = tmp_path / "request.json"
    atomic_json(request, {"engine": "diarization", "model": str(tmp_path / "missing"), "audio": str(tmp_path / "missing.wav")})
    with pytest.raises(RuntimeError, match="FileNotFoundError"):
        run_worker(request, timeout=10)

def test_worker_timeout_is_bounded(tmp_path):
    request = tmp_path / "request.json"
    atomic_json(request, {"engine": "diarization", "model": str(tmp_path / "missing")})
    with pytest.raises(TimeoutError, match="WORKER_TIMEOUT"):
        run_worker(request, timeout=0.001)
