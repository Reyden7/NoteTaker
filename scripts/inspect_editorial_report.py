"""Compare immutable source data and draft document properties after a real run."""
import hashlib
import json
import shutil
import sqlite3
from pathlib import Path
from notetaker.services.storage import SessionStore
from notetaker.services.session_export import export_session
from notetaker.services.voice_profiles import VoiceProfiles
from pypdf import PdfReader
from docx import Document

repo=Path(__file__).resolve().parents[1]
root=Path((repo/'output/editorial-report-latest.txt').read_text().strip())
previous=Path((repo/'output/automatic-report-latest.txt').read_text().strip())
run=root/'municipal-run-04';working=run/'session-copy'
original=previous/'municipal-v3/session-copy'
sid='58709d31-6b8f-434e-ac75-73871d745298'

def connect(path):return sqlite3.connect((path/'notetaker.sqlite3').resolve().as_uri()+'?mode=ro',uri=True)
unchanged={}
with connect(original) as a,connect(working) as b:
    for table, in a.execute("SELECT name FROM sqlite_master WHERE type='table'"):
        if table.startswith('sqlite_') or table in ('job','draft_minutes'):continue
        query='SELECT * FROM "'+table+'"'
        unchanged[table]=sorted(a.execute(query).fetchall(),key=repr)==sorted(b.execute(query).fetchall(),key=repr)
    olds=a.execute('SELECT * FROM draft_minutes').fetchall()
    news=b.execute('SELECT * FROM draft_minutes').fetchall()
    unchanged['prior_drafts']=all(r in news for r in olds)
    unchanged['no_new_review']=all(r[4] is None for r in news if r not in olds)
assert all(unchanged.values()),unchanged
store=SessionStore(working);state=store.get(sid)
assert state['draft']['review'] is None
try:export_session(store,sid,root/'OFFICIAL-EXPORT-MUST-NOT-EXIST.pdf')
except ValueError:blocked=True
else:blocked=False
assert blocked
result=json.loads((root/'livrables-v4/result.json').read_text(encoding='utf-8'))
job=Path(result['workspace'])
assert job.resolve().is_relative_to((working/'sessions'/sid/'jobs').resolve())
for name in ('verification-report.md','automatic-inventory.json','editorial-plan.json','reconstructed-turns.json'):
    shutil.copyfile(job/name,root/name)
store.close()
vault=VoiceProfiles(working)
assert not vault.policy().get('enabled') and not vault.samples()

def properties(directory):
    pdf=directory/'Compte-rendu-BROUILLON.pdf'
    if not pdf.exists():pdf=directory/'compte-rendu-BROUILLON.pdf'
    docx=directory/'Compte-rendu-BROUILLON.docx'
    if not docx.exists():docx=directory/'compte-rendu-BROUILLON.docx'
    reader=PdfReader(pdf);text='\n'.join(p.extract_text() for p in reader.pages)
    document=Document(docx)
    body='\n'.join(p.text for p in document.paragraphs)
    return {'pdf_pages':len(reader.pages),'generic_la_personne':body.casefold().count('la personne'),
        'quote_openings':body.count('«'),'technical_diagnostic_in_document':'Diagnostic :' in body,
        'docx_paragraphs':len(document.paragraphs),'pdf_words':len(text.split()),
        'pdf_draft_mark_on_every_page':all('BROUILLON NON RELU' in p.extract_text() for p in reader.pages),
        'docx_draft_footer':'BROUILLON NON RELU' in document.sections[0].footer.paragraphs[0].text,
        'pdf_sha256':hashlib.sha256(pdf.read_bytes()).hexdigest(),'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest()}

old=properties(previous/'delivery-final');new=properties(root/'livrables-v4')
assert new['pdf_draft_mark_on_every_page'] and new['docx_draft_footer']
assert new['generic_la_personne']==0 and not new['technical_diagnostic_in_document']
result={'preservation':unchanged,'official_export_blocked':blocked,'voice_identification_disabled':True,
    'old_document':old,'new_document':new,'human_fidelity_measured':False,'human_corrections_measured':False}
(root/'delivery-checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
(root/'new-draft-readable.json').write_text(json.dumps(state['draft']['content'],ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
